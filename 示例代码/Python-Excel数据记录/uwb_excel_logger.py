#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
============================================================================
UWB单基站跟随套件 - 球坐标转三维坐标并存储到Excel
============================================================================

功能说明：
    本程序从UWB基站接收串口数据，将球坐标（距离、方位角、仰角）
    转换为三维坐标（X、Y、Z），并实时保存到Excel文件中。

硬件连接：
    - UWB基站通过USB转TTL模块连接到电脑
    - 串口参数：115200波特率，8数据位，1停止位，无校验

依赖安装：
    pip install pyserial openpyxl

使用方法：
    python uwb_excel_logger.py [串口号] [输出文件名]
    例如：python uwb_excel_logger.py COM3 uwb_data.xlsx
    
作者：Copilot
日期：2026-01-22
============================================================================
"""

import sys
import math
import struct
import threading
import time
import os
from datetime import datetime
from collections import deque

# 尝试导入依赖库
try:
    import serial
    import serial.tools.list_ports
except ImportError:
    print("错误：请先安装 pyserial 库")
    print("运行命令：pip install pyserial")
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("错误：请先安装 openpyxl 库")
    print("运行命令：pip install openpyxl")
    sys.exit(1)

# ============================================================================
# 全局配置
# ============================================================================

# 协议相关常量
FRAME_HEADER = b'\xff\xff\xff\xff'  # 帧头
FRAME_LENGTH = 37                    # 数据帧长度（字节）
CMD_POSITION = 0x2001                # 位置数据命令字

# 保存配置
SAVE_INTERVAL = 10          # 每隔多少条数据自动保存一次
DEFAULT_FILENAME = "uwb_data.xlsx"

# ============================================================================
# 数据解析类
# ============================================================================

class UWBDataParser:
    """
    UWB数据帧解析器
    
    协议格式（共37字节）：
    - MessageHeader:   0-3   (4字节)  帧头 0xFFFFFFFF
    - PacketLength:    4-5   (2字节)  消息体长度
    - SequenceID:      6-7   (2字节)  消息流水号
    - RequestCommand:  8-9   (2字节)  命令码 0x2001
    - VersionID:       10-11 (2字节)  协议版本 0x0100
    - AnchorID:        12-15 (4字节)  基站ID
    - TagID:           16-19 (4字节)  标签ID
    - Distance:        20-23 (4字节)  距离，单位：cm
    - Azimuth:         24-25 (2字节)  方位角，单位：度（有符号）
    - Elevation:       26-27 (2字节)  仰角，单位：度（有符号）
    - TagStatus:       28-29 (2字节)  标签状态
    - BatchSn:         30-31 (2字节)  测距序号
    - Reserve:         32-35 (4字节)  预留
    - XorByte:         36    (1字节)  异或校验
    """
    
    def __init__(self):
        self.buffer = b''
    
    def calculate_xor(self, data):
        """计算异或校验值"""
        xor_value = 0
        for byte in data:
            xor_value ^= byte
        return xor_value
    
    def parse(self, new_data):
        """
        解析接收到的数据
        
        参数：
            new_data: 新接收的字节数据
            
        返回：
            解析成功返回字典列表，包含距离、角度等信息
            解析失败返回 None
        """
        self.buffer += new_data
        results = []
        
        while True:
            # 查找帧头
            header_idx = self.buffer.find(FRAME_HEADER)
            
            if header_idx < 0:
                # 未找到帧头，保留最后3字节（可能是不完整的帧头）
                if len(self.buffer) > 3:
                    self.buffer = self.buffer[-3:]
                break
            
            # 移除帧头之前的无效数据
            if header_idx > 0:
                self.buffer = self.buffer[header_idx:]
            
            # 检查数据长度是否足够
            if len(self.buffer) < FRAME_LENGTH:
                break
            
            # 提取一帧数据
            frame = self.buffer[:FRAME_LENGTH]
            
            # 验证异或校验
            calculated_xor = self.calculate_xor(frame[:-1])
            received_xor = frame[-1]
            
            if calculated_xor != received_xor:
                # 校验失败，跳过这个帧头，继续搜索
                self.buffer = self.buffer[4:]
                continue
            
            # 解析数据（大端序）
            try:
                command = struct.unpack('>H', frame[8:10])[0]
                
                if command == CMD_POSITION:
                    # 解析各字段（修正字节偏移）
                    anchor_id = struct.unpack('>I', frame[12:16])[0]
                    tag_id = struct.unpack('>I', frame[16:20])[0]
                    distance_cm = struct.unpack('>I', frame[20:24])[0]
                    azimuth_deg = struct.unpack('>h', frame[24:26])[0]  # 有符号
                    elevation_deg = struct.unpack('>h', frame[26:28])[0]  # 有符号
                    
                    # 将球坐标转换为笛卡尔坐标
                    x, y, z = self.spherical_to_cartesian(
                        distance_cm, azimuth_deg, elevation_deg
                    )
                    
                    result = {
                        'timestamp': datetime.now(),
                        'anchor_id': anchor_id,
                        'tag_id': tag_id,
                        'distance_cm': distance_cm,
                        'azimuth_deg': azimuth_deg,
                        'elevation_deg': elevation_deg,
                        'x': x,
                        'y': y,
                        'z': z
                    }
                    results.append(result)
            except Exception as e:
                print(f"解析错误: {e}")
            
            # 移除已处理的数据
            self.buffer = self.buffer[FRAME_LENGTH:]
        
        return results if results else None
    
    def spherical_to_cartesian(self, distance_cm, azimuth_deg, elevation_deg):
        """
        将球坐标转换为笛卡尔坐标
        
        参数：
            distance_cm:   距离（厘米）
            azimuth_deg:   方位角（度），正前方为0度，顺时针为正
            elevation_deg: 俯仰角（度），水平为0度，向上为正
            
        返回：
            (x, y, z) 坐标，单位：厘米
            
        坐标系定义：
            - Y轴：正前方
            - X轴：右侧为正
            - Z轴：向上为正
        """
        # 转换为弧度
        azimuth_rad = math.radians(azimuth_deg)
        elevation_rad = math.radians(elevation_deg)
        
        # 计算水平投影距离
        horizontal_distance = distance_cm * math.cos(elevation_rad)
        
        # 计算XYZ坐标
        x = horizontal_distance * math.sin(azimuth_rad)
        y = horizontal_distance * math.cos(azimuth_rad)
        z = distance_cm * math.sin(elevation_rad)
        
        return round(x, 2), round(y, 2), round(z, 2)


# ============================================================================
# Excel记录器类
# ============================================================================

class ExcelLogger:
    """Excel数据记录器"""
    
    def __init__(self, filename):
        """
        初始化Excel记录器
        
        参数：
            filename: Excel文件名
        """
        self.filename = filename
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "UWB定位数据"
        self.row_count = 1
        self.data_count = 0
        
        # 设置表头
        self._setup_header()
    
    def _setup_header(self):
        """设置Excel表头"""
        headers = [
            "序号",
            "时间戳",
            "基站ID",
            "标签ID",
            "距离(cm)",
            "方位角(°)",
            "仰角(°)",
            "X坐标(cm)",
            "Y坐标(cm)",
            "Z坐标(cm)"
        ]
        
        # 设置表头样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        column_widths = [8, 22, 12, 12, 12, 12, 12, 12, 12, 12]
        for col, width in enumerate(column_widths, 1):
            self.worksheet.column_dimensions[chr(64 + col)].width = width
        
        self.row_count = 2  # 数据从第2行开始
    
    def add_record(self, data):
        """
        添加一条数据记录
        
        参数：
            data: 包含定位数据的字典
        """
        self.data_count += 1
        
        # 准备数据
        row_data = [
            self.data_count,
            data['timestamp'].strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
            f"0x{data['anchor_id']:08X}",
            f"0x{data['tag_id']:08X}",
            data['distance_cm'],
            data['azimuth_deg'],
            data['elevation_deg'],
            data['x'],
            data['y'],
            data['z']
        ]
        
        # 设置单元格样式
        alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 写入数据
        for col, value in enumerate(row_data, 1):
            cell = self.worksheet.cell(row=self.row_count, column=col, value=value)
            cell.alignment = alignment
            cell.border = thin_border
        
        self.row_count += 1
    
    def save(self):
        """保存Excel文件"""
        try:
            self.workbook.save(self.filename)
            return True
        except Exception as e:
            print(f"保存文件失败: {e}")
            return False
    
    def close(self):
        """关闭工作簿"""
        try:
            self.workbook.close()
        except:
            pass


# ============================================================================
# 串口通信类
# ============================================================================

class SerialReceiver:
    """串口数据接收器"""
    
    def __init__(self, port, baudrate=115200):
        """
        初始化串口接收器
        
        参数：
            port: 串口号，如 'COM3' 或 '/dev/ttyUSB0'
            baudrate: 波特率，默认115200
        """
        self.port = port
        self.baudrate = baudrate
        self.serial = None
        self.running = False
        self.parser = UWBDataParser()
        self.data_queue = deque(maxlen=1000)
        self.thread = None
    
    def connect(self):
        """连接串口"""
        try:
            self.serial = serial.Serial(
                port=self.port,
                baudrate=self.baudrate,
                bytesize=serial.EIGHTBITS,
                stopbits=serial.STOPBITS_ONE,
                parity=serial.PARITY_NONE,
                timeout=0.1
            )
            print(f"✓ 串口 {self.port} 已连接")
            print(f"  波特率: {self.baudrate}")
            return True
        except Exception as e:
            print(f"✗ 串口连接失败: {e}")
            return False
    
    def disconnect(self):
        """断开串口连接"""
        self.running = False
        if self.thread:
            self.thread.join(timeout=1)
        if self.serial and self.serial.is_open:
            self.serial.close()
            print("串口已断开")
    
    def start_receiving(self):
        """开始接收数据（在新线程中）"""
        if not self.serial or not self.serial.is_open:
            print("错误：串口未连接")
            return False
        
        self.running = True
        self.thread = threading.Thread(target=self._receive_loop, daemon=True)
        self.thread.start()
        print("✓ 开始接收数据...")
        return True
    
    def _receive_loop(self):
        """数据接收循环（在后台线程中运行）"""
        while self.running:
            try:
                if self.serial.in_waiting > 0:
                    data = self.serial.read(self.serial.in_waiting)
                    results = self.parser.parse(data)
                    if results:
                        for result in results:
                            self.data_queue.append(result)
            except Exception as e:
                print(f"接收错误: {e}")
                time.sleep(0.1)
    
    def get_all_data(self):
        """获取所有待处理的数据"""
        data_list = list(self.data_queue)
        self.data_queue.clear()
        return data_list


# ============================================================================
# 模拟数据生成器（用于测试）
# ============================================================================

class SimulatedReceiver:
    """模拟数据接收器（用于无硬件时测试）"""
    
    def __init__(self):
        self.data_queue = deque(maxlen=100)
        self.running = False
        self.thread = None
        self.angle = 0
        self.radius = 50
        self.z_angle = 0
        import random
        self.random = random
    
    def connect(self):
        print("✓ 模拟模式已启动（无需真实硬件）")
        return True
    
    def disconnect(self):
        self.running = False
        if self.thread:
            self.thread.join(timeout=1)
    
    def start_receiving(self):
        self.running = True
        self.thread = threading.Thread(target=self._generate_loop, daemon=True)
        self.thread.start()
        print("✓ 开始生成模拟数据...")
        return True
    
    def _generate_loop(self):
        """生成模拟数据"""
        while self.running:
            # 生成螺旋线轨迹
            self.angle += 5
            self.radius += 0.2
            self.z_angle += 2
            
            # 计算模拟的球坐标
            distance = int(self.radius + self.random.uniform(-2, 2))
            azimuth = int(self.angle % 360 - 180 + self.random.uniform(-2, 2))
            elevation = int(15 * math.sin(math.radians(self.z_angle)) + self.random.uniform(-1, 1))
            
            # 计算XYZ坐标
            azimuth_rad = math.radians(azimuth)
            elevation_rad = math.radians(elevation)
            horizontal_distance = distance * math.cos(elevation_rad)
            x = round(horizontal_distance * math.sin(azimuth_rad), 2)
            y = round(horizontal_distance * math.cos(azimuth_rad), 2)
            z = round(distance * math.sin(elevation_rad), 2)
            
            data = {
                'timestamp': datetime.now(),
                'anchor_id': 0xAAA2,
                'tag_id': 0xAAA1,
                'distance_cm': distance,
                'azimuth_deg': azimuth,
                'elevation_deg': elevation,
                'x': x,
                'y': y,
                'z': z
            }
            self.data_queue.append(data)
            
            time.sleep(0.2)  # 5Hz 数据率
    
    def get_all_data(self):
        data_list = list(self.data_queue)
        self.data_queue.clear()
        return data_list


# ============================================================================
# 辅助函数
# ============================================================================

def list_serial_ports():
    """列出所有可用的串口"""
    ports = serial.tools.list_ports.comports()
    if not ports:
        print("未检测到可用串口")
        return []
    
    print("\n可用串口列表：")
    print("-" * 50)
    for i, port in enumerate(ports, 1):
        print(f"  {i}. {port.device} - {port.description}")
    print("-" * 50)
    
    return [port.device for port in ports]


def print_usage():
    """打印使用说明"""
    print("""
╔══════════════════════════════════════════════════════════════╗
║      UWB单基站跟随套件 - 球坐标转三维坐标Excel记录程序       ║
╠══════════════════════════════════════════════════════════════╣
║                                                              ║
║  用法：python uwb_excel_logger.py [选项]                     ║
║                                                              ║
║  选项：                                                      ║
║    <串口号>           指定串口，如 COM3 或 /dev/ttyUSB0      ║
║    <串口号> <文件名>  指定串口和输出文件名                   ║
║    --list             列出所有可用串口                       ║
║    --simulate         使用模拟数据（无需硬件）               ║
║    --help             显示此帮助信息                         ║
║                                                              ║
║  示例：                                                      ║
║    python uwb_excel_logger.py COM3                           ║
║    python uwb_excel_logger.py COM3 my_data.xlsx              ║
║    python uwb_excel_logger.py --simulate                     ║
║    python uwb_excel_logger.py --simulate test.xlsx           ║
║                                                              ║
║  依赖安装：                                                  ║
║    pip install pyserial openpyxl                             ║
║                                                              ║
║  输出：                                                      ║
║    Excel文件包含：序号、时间戳、基站ID、标签ID、             ║
║    距离(cm)、方位角(°)、仰角(°)、X/Y/Z坐标(cm)               ║
║                                                              ║
╚══════════════════════════════════════════════════════════════╝
""")


def print_data_row(data, count):
    """打印一行数据到控制台"""
    print(f"[{count:5d}] 距离:{data['distance_cm']:4d}cm  "
          f"方位角:{data['azimuth_deg']:4d}°  仰角:{data['elevation_deg']:3d}°  →  "
          f"X:{data['x']:8.2f}  Y:{data['y']:8.2f}  Z:{data['z']:8.2f} cm")


# ============================================================================
# 主函数
# ============================================================================

def main():
    """主函数"""
    # 解析命令行参数
    args = sys.argv[1:]
    
    if len(args) == 0:
        print_usage()
        ports = list_serial_ports()
        if ports:
            print("\n提示：请指定串口号，例如：")
            print(f"  python {sys.argv[0]} {ports[0]}")
            print(f"\n或使用模拟模式测试：")
            print(f"  python {sys.argv[0]} --simulate")
        return
    
    # 帮助信息
    if args[0] in ['--help', '-h']:
        print_usage()
        return
    
    # 列出串口
    if args[0] == '--list':
        list_serial_ports()
        return
    
    # 确定模式和文件名
    simulate_mode = False
    filename = DEFAULT_FILENAME
    port = None
    
    if args[0] == '--simulate':
        simulate_mode = True
        if len(args) > 1:
            filename = args[1]
    else:
        port = args[0]
        if len(args) > 1:
            filename = args[1]
    
    # 确保文件名以.xlsx结尾
    if not filename.endswith('.xlsx'):
        filename += '.xlsx'
    
    print("\n" + "="*60)
    print("    UWB 球坐标 → 三维坐标 Excel记录器")
    print("="*60)
    
    # 创建接收器
    if simulate_mode:
        print("\n模式: 模拟数据")
        receiver = SimulatedReceiver()
    else:
        print(f"\n串口: {port}")
        receiver = SerialReceiver(port)
    
    print(f"输出文件: {filename}")
    
    # 连接
    if not receiver.connect():
        return
    
    # 开始接收数据
    if not receiver.start_receiving():
        receiver.disconnect()
        return
    
    # 创建Excel记录器
    logger = ExcelLogger(filename)
    print(f"\n✓ Excel文件已创建: {filename}")
    
    print("\n" + "-"*60)
    print("开始记录数据... (按 Ctrl+C 停止)")
    print("-"*60)
    print("\n[序号 ] 距离        方位角      仰角        →  X坐标     Y坐标     Z坐标")
    print("-"*80)
    
    data_count = 0
    last_save_count = 0
    
    try:
        while True:
            # 获取新数据
            data_list = receiver.get_all_data()
            
            for data in data_list:
                data_count += 1
                
                # 添加到Excel
                logger.add_record(data)
                
                # 打印到控制台
                print_data_row(data, data_count)
                
                # 定期保存
                if data_count - last_save_count >= SAVE_INTERVAL:
                    logger.save()
                    last_save_count = data_count
                    print(f"  [已自动保存 {data_count} 条数据]")
            
            time.sleep(0.05)
            
    except KeyboardInterrupt:
        print("\n\n" + "-"*60)
        print("停止记录...")
    finally:
        # 最终保存
        logger.save()
        logger.close()
        receiver.disconnect()
        
        print(f"\n✓ 数据已保存到: {os.path.abspath(filename)}")
        print(f"✓ 共记录 {data_count} 条数据")
        print("\n程序已退出")


if __name__ == '__main__':
    main()
